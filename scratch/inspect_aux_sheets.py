import pandas as pd
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
excel_path = 'Desempeño Fiscal departamentos 2020-2024.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=False) # Load with formulas to see if they are formulas or static!

print("=== VERIFICACION DE FORMULAS VS VALORES EN HOJAS ===")

sheets_to_check = ['Dashboard IDF', 'Ejemplo Cálculo IDF', 'Pivot prom 2020-2024', 'Tendencias', 'Gráficas', 'Mejores-Peores', 'Análisis de Brechas', 'Analisis rangos desempeño']

for s in sheets_to_check:
    ws = wb[s]
    # Check sample formulas
    formulas = []
    for r in range(1, min(ws.max_row+1, 30)):
        for c in range(1, min(ws.max_column+1, 20)):
            val = ws.cell(row=r, column=c).value
            if str(val).startswith('='):
                formulas.append((r, c, val))
    print(f"\nHoja '{s}': {len(formulas)} fórmulas encontradas en las primeras filas/columnas.")
    if formulas:
        print("  Ejemplos de fórmulas:", formulas[:5])

# Now load data_only=True to inspect the contents and purpose of these sheets
wb_vals = openpyxl.load_workbook(excel_path, data_only=True)

print("\n=== CONTENIDO DETALLADO DE LAS OTRAS HOJAS ===")

# 1. Pivot prom 2020-2024
print("\n--- 1. Pivot prom 2020-2024 ---")
df_pivot = pd.read_excel(excel_path, sheet_name='Pivot prom 2020-2024')
print(df_pivot.to_string())

# 2. Tendencias
print("\n--- 2. Tendencias ---")
df_tend = pd.read_excel(excel_path, sheet_name='Tendencias', header=None).dropna(how='all')
print("Shape:", df_tend.shape)
print(df_tend.head(15).to_string())

# 3. Mejores-Peores
print("\n--- 3. Mejores-Peores ---")
df_mp = pd.read_excel(excel_path, sheet_name='Mejores-Peores', header=None).dropna(how='all')
print("Shape:", df_mp.shape)
print(df_mp.head(20).to_string())

# 4. Análisis de Brechas
print("\n--- 4. Análisis de Brechas ---")
df_brechas = pd.read_excel(excel_path, sheet_name='Análisis de Brechas', header=None).dropna(how='all')
print("Shape:", df_brechas.shape)
print(df_brechas.head(15).to_string())

# 5. Analisis rangos desempeño
print("\n--- 5. Analisis rangos desempeño ---")
df_rangos = pd.read_excel(excel_path, sheet_name='Analisis rangos desempeño', header=None).dropna(how='all')
print("Shape:", df_rangos.shape)
print(df_rangos.to_string())

# 6. Gráficas
print("\n--- 6. Gráficas ---")
df_graf = pd.read_excel(excel_path, sheet_name='Gráficas', header=None).dropna(how='all')
print("Shape:", df_graf.shape)
print(df_graf.head(15).to_string())

# 7. Dashboard IDF
print("\n--- 7. Dashboard IDF ---")
df_dash = pd.read_excel(excel_path, sheet_name='Dashboard IDF', header=None)
print("Dashboard IDF rows with data:")
for r in range(len(df_dash)):
    row = df_dash.iloc[r].dropna()
    if len(row) > 0:
        row_str = " | ".join([f"C{c}:{str(val)[:20]}" for c, val in row.items()])
        print(f"R{r:02d}: {row_str[:120]}")

# 8. Ejemplo Cálculo IDF
print("\n--- 8. Ejemplo Cálculo IDF ---")
df_ej = pd.read_excel(excel_path, sheet_name='Ejemplo Cálculo IDF', header=None)
for r in range(len(df_ej)):
    row = df_ej.iloc[r].dropna()
    if len(row) > 0:
        row_str = " | ".join([f"C{c}:{str(val)[:30]}" for c, val in row.items()])
        print(f"R{r:02d}: {row_str[:120]}")
